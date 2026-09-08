"""
Markdown → Word (.docx) Exporter
Converts any .md file to a premium-styled, export-ready .docx document.
Part of the markdown-docx-exporter skill.

Usage:
    python convert.py <input.md> <output.docx> [--title "Title"]

Uses the Bluejaypro Deployment RAG default theme.
"""

import os
import sys
import re
import argparse
import markdown
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


# ─── Theme Definitions ───────────────────────────────────────────────────────

THEME = {
    "heading_colors": {
        1: RGBColor(31, 78, 120),
        2: RGBColor(46, 117, 182),
        3: RGBColor(65, 113, 156),
    },
    "heading_default": RGBColor(51, 51, 51),
    "table_header_bg": "1F4E78",
    "table_header_fg": RGBColor(255, 255, 255),
    "table_alt_row": "F2F5F8",
    "table_even_row": "FFFFFF",
    "code_block_bg": "F4F4F5",
    "code_block_fg": RGBColor(30, 30, 30),
    "inline_code_fg": RGBColor(199, 37, 78),
    "link_color": RGBColor(31, 78, 120),
    "blockquote_bg": "F9FAFB",
    "blockquote_fg": RGBColor(100, 110, 120),
    "body_fg": RGBColor(51, 51, 51),
    "hr_color": "CCCCCC",
    "header_fg": RGBColor(128, 128, 128),
}


# ─── Word XML Helpers ─────────────────────────────────────────────────────────

def set_cell_shading(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_paragraph_shading(p, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    p._p.get_or_add_pPr().append(parse_xml(shading_xml))

def set_paragraph_bottom_border(p, color_hex="D3D3D3"):
    pBdr_xml = (
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="{color_hex}"/>'
        f'</w:pBdr>'
    )
    p._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

def set_paragraph_left_border(p, color_hex="2E75B6"):
    pBdr_xml = (
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="18" w:space="8" w:color="{color_hex}"/>'
        f'</w:pBdr>'
    )
    p._p.get_or_add_pPr().append(parse_xml(pBdr_xml))

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(
        r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')
    )
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)


# ─── Auto-Title Detection ────────────────────────────────────────────────────

def detect_title(md_text, filepath):
    """Extract the first H1 heading as the document title, or fall back to filename."""
    match = re.search(r'^#\s+(.+)$', md_text, re.MULTILINE)
    if match:
        # Strip any markdown formatting from the title
        title = match.group(1).strip()
        title = re.sub(r'[*_`]', '', title)
        return title
    # Fallback: use the filename without extension
    basename = os.path.splitext(os.path.basename(filepath))[0]
    return basename.replace('-', ' ').replace('_', ' ').title()


# ─── Inline Element Processing ───────────────────────────────────────────────

def process_inline_element(p, element, theme, bold=False, italic=False,
                           is_code=False, is_link=False, href=""):
    if isinstance(element, str):
        text = element
        if is_link and href:
            run = p.add_run(
                f"{text} ({href})"
                if text.strip() != href.strip() and not href.startswith('#')
                else text
            )
        else:
            run = p.add_run(text)

        run.bold = bold
        run.italic = italic
        run.font.name = 'Segoe UI'
        if is_code:
            run.font.name = 'Consolas'
            run.font.size = Pt(9.0)
            run.font.color.rgb = theme["inline_code_fg"]
        if is_link:
            run.font.color.rgb = theme["link_color"]
            run.underline = True
        return

    for child in element.children:
        if isinstance(child, str):
            text = child
            if is_link and href:
                run = p.add_run(
                    f"{text} ({href})"
                    if text.strip() != href.strip() and not href.startswith('#')
                    else text
                )
            else:
                run = p.add_run(text)

            run.bold = bold
            run.italic = italic
            run.font.name = 'Segoe UI'
            if is_code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9.0)
                run.font.color.rgb = theme["inline_code_fg"]
            if is_link:
                run.font.color.rgb = theme["link_color"]
                run.underline = True
        else:
            c_bold = bold or (child.name in ('strong', 'b'))
            c_italic = italic or (child.name in ('em', 'i'))
            c_code = is_code or (child.name == 'code')
            c_link = is_link or (child.name == 'a')
            c_href = href if child.name != 'a' else child.get('href', '')
            process_inline_element(
                p, child, theme,
                bold=c_bold, italic=c_italic,
                is_code=c_code, is_link=c_link, href=c_href
            )


# ─── Block-Level Processing ──────────────────────────────────────────────────

def add_heading_with_styling(doc, text_element, level, theme):
    p = doc.add_heading(level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level > 1 else 22)
    p.paragraph_format.space_after = Pt(6)

    process_inline_element(p, text_element, theme)

    heading_color = theme["heading_colors"].get(level, theme["heading_default"])
    size_map = {1: Pt(18), 2: Pt(14), 3: Pt(12.5), 4: Pt(11.5)}
    heading_size = size_map.get(level, Pt(11))

    for run in p.runs:
        run.font.name = 'Segoe UI'
        run.font.color.rgb = heading_color
        run.font.size = heading_size
        run.bold = True

    return p


def process_list_item(doc, li_element, depth, list_style_base, theme):
    inline_nodes = []
    nested_lists = []
    for child in li_element.children:
        if isinstance(child, str):
            inline_nodes.append(child)
        elif child.name in ('ul', 'ol'):
            nested_lists.append(child)
        else:
            inline_nodes.append(child)

    style_name = list_style_base
    if depth > 1:
        style_name = f"{list_style_base} {depth}"

    try:
        p = doc.add_paragraph(style=style_name)
    except KeyError:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25 * depth)
        if 'Bullet' in list_style_base:
            p.add_run("•  ")
        else:
            p.add_run("-  ")

    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15

    for node in inline_nodes:
        if isinstance(node, str):
            run = p.add_run(node)
            run.font.name = 'Segoe UI'
        else:
            process_inline_element(p, node, theme)

    for nested in nested_lists:
        if nested.name == 'ul':
            process_list(doc, nested, depth + 1, 'List Bullet', theme)
        elif nested.name == 'ol':
            process_list(doc, nested, depth + 1, 'List Number', theme)


def process_list(doc, list_element, depth, list_style_base, theme):
    for child in list_element.children:
        if isinstance(child, str):
            continue
        if child.name == 'li':
            process_list_item(doc, child, depth, list_style_base, theme)


def process_table(doc, table_element, theme):
    rows = table_element.find_all('tr')
    if not rows:
        return

    cols = rows[0].find_all(['td', 'th'])
    num_cols = len(cols)
    num_rows = len(rows)

    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = 'Table Grid'

    for r_idx, row in enumerate(rows):
        cells = row.find_all(['td', 'th'])
        for c_idx, cell in enumerate(cells):
            if c_idx >= num_cols:
                break
            docx_cell = table.cell(r_idx, c_idx)
            p = docx_cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15

            is_header = (cell.name == 'th' or r_idx == 0)
            process_inline_element(p, cell, theme, bold=is_header)

            if is_header:
                set_cell_shading(docx_cell, theme["table_header_bg"])
                for run in p.runs:
                    run.font.color.rgb = theme["table_header_fg"]
                    run.bold = True
            else:
                if r_idx % 2 == 1:
                    set_cell_shading(docx_cell, theme["table_alt_row"])
                else:
                    set_cell_shading(docx_cell, theme["table_even_row"])


def process_blockquote(doc, bq_element, theme):
    """Handle blockquotes with left-border accent."""
    for child in bq_element.children:
        if isinstance(child, str):
            text = child.strip()
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.space_after = Pt(6)
                set_paragraph_shading(p, theme["blockquote_bg"])
                # Grab the primary heading color for the left border accent
                accent = theme["heading_colors"].get(2, theme["heading_default"])
                accent_hex = f"{accent[0]:02X}{accent[1]:02X}{accent[2]:02X}"
                set_paragraph_left_border(p, accent_hex)
                run = p.add_run(text)
                run.font.name = 'Segoe UI'
                run.italic = True
                run.font.color.rgb = theme["blockquote_fg"]
        elif child.name == 'p':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            set_paragraph_shading(p, theme["blockquote_bg"])
            accent = theme["heading_colors"].get(2, theme["heading_default"])
            accent_hex = f"{accent[0]:02X}{accent[1]:02X}{accent[2]:02X}"
            set_paragraph_left_border(p, accent_hex)
            process_inline_element(p, child, theme, italic=True)
            for run in p.runs:
                run.font.color.rgb = theme["blockquote_fg"]
        elif child.name == 'blockquote':
            # Nested blockquote
            process_blockquote(doc, child, theme)
        elif child.name in ('ul', 'ol'):
            style = 'List Bullet' if child.name == 'ul' else 'List Number'
            process_list(doc, child, 1, style, theme)


# ─── Main Converter ──────────────────────────────────────────────────────────

def convert(input_path, output_path, title=None):
    """Convert a markdown file to a styled .docx document."""
    theme = THEME

    if not os.path.exists(input_path):
        print(f"ERROR: Input file '{input_path}' not found.")
        sys.exit(1)

    print(f"[1/5] Reading: {input_path}")
    with open(input_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    # Auto-detect title if not provided
    if not title:
        title = detect_title(md_text, input_path)
    print(f"[2/5] Title: {title}")

    print("[3/5] Converting Markdown -> HTML...")
    html = markdown.markdown(
        md_text,
        extensions=['tables', 'fenced_code', 'nl2br']
    )
    soup = BeautifulSoup(html, 'html.parser')

    print("[4/5] Building Word document...")
    doc = Document()

    # ── Page Setup ──
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(title)
        hrun.font.name = 'Segoe UI'
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = theme["header_fg"]

        # Footer with page number
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Page ")
        frun.font.name = 'Segoe UI'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = theme["header_fg"]
        add_page_number(frun)

    # ── Default Style ──
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Segoe UI'
    font.size = Pt(10.5)
    font.color.rgb = theme["body_fg"]

    # ── Process All Blocks ──
    for element in soup.children:
        if isinstance(element, str):
            continue

        tag = element.name

        # Headings
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
            level = int(tag[1])
            add_heading_with_styling(doc, element, level, theme)

        # Paragraphs
        elif tag == 'p':
            img_tag = element.find('img')
            if img_tag:
                raw_src = img_tag.get('src', '')
                src = raw_src.replace('file:///', '').replace('file://', '')
                if not os.path.isabs(src):
                    src = os.path.normpath(os.path.join(os.path.dirname(input_path), src))
                if os.path.exists(src):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run()
                    try:
                        run.add_picture(src, width=Inches(6.0))
                    except Exception as img_err:
                        print(f"Warning: Could not insert image {src}: {img_err}")
                    alt = img_tag.get('alt')
                    if alt:
                        cp = doc.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(8)
                        crun = cp.add_run(f"Figure: {alt}")
                        crun.font.name = 'Segoe UI'
                        crun.font.size = Pt(8.5)
                        crun.font.italic = True
                        crun.font.color.rgb = theme["header_fg"]
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.15
                    process_inline_element(p, element, theme)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                process_inline_element(p, element, theme)

        # Blockquotes
        elif tag == 'blockquote':
            process_blockquote(doc, element, theme)

        # Code blocks
        elif tag == 'pre':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            set_paragraph_shading(p, theme["code_block_bg"])

            code_text = element.get_text()
            run = p.add_run(code_text)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = theme["code_block_fg"]

        # Lists
        elif tag == 'ul':
            process_list(doc, element, 1, 'List Bullet', theme)
        elif tag == 'ol':
            process_list(doc, element, 1, 'List Number', theme)

        # Tables
        elif tag == 'table':
            process_table(doc, element, theme)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)

        # Horizontal rules
        elif tag == 'hr':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            set_paragraph_bottom_border(p, theme["hr_color"])

    # ── Save ──
    print(f"[5/5] Saving: {output_path}")
    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    doc.save(output_path)

    size_kb = os.path.getsize(output_path) / 1024
    print(f"Done! ({size_kb:.1f} KB)")
    return output_path


# ─── CLI Entry Point ─────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Convert Markdown to a premium-styled Word document (.docx)"
    )
    parser.add_argument("input", help="Path to the input .md file")
    parser.add_argument("output", help="Path for the output .docx file")
    parser.add_argument(
        "--title", default=None,
        help="Document title for the header (auto-detected from H1 if omitted)"
    )
    args = parser.parse_args()
    convert(args.input, args.output, title=args.title)


if __name__ == '__main__':
    main()
